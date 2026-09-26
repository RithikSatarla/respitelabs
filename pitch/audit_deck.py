"""Check the built deck for the things that make a slide look broken.

Catches shapes that run off the slide, text boxes too small for the text they
hold, text under the minimum size, and text boxes that overlap each other.
It is an estimate, not a renderer, so treat hits as "go and look at that slide".

Run from the repo root:  python pitch/audit_deck.py
"""

import pathlib

from pptx import Presentation
from pptx.util import Emu

ROOT = pathlib.Path(__file__).resolve().parent.parent
DECK = ROOT / "pitch" / "Kintrace-deck.pptx"

MIN_PT = 9.0          # anything smaller than this is unreadable on a deck
EMU_IN = 914400.0


def inches(v):
    return (v or 0) / EMU_IN


def est_lines(text, width_in, size_pt):
    """Rough wrap estimate. Arial averages about 0.5 em per character."""
    if not text.strip():
        return 0
    per_line = max(1, int(width_in * 72 / (size_pt * 0.505)))
    lines = 0
    for para in text.split("\n"):
        lines += max(1, -(-len(para) // per_line))
    return lines


def main():
    prs = Presentation(str(DECK))
    SW, SH = inches(prs.slide_width), inches(prs.slide_height)
    print(f"slide {SW:.3f} x {SH:.3f} in, {len(prs.slides._sldIdLst)} slides\n")
    problems = 0

    for n, slide in enumerate(prs.slides, start=1):
        hits = []
        boxes = []
        for sh in slide.shapes:
            x, y = inches(sh.left), inches(sh.top)
            w, h = inches(sh.width), inches(sh.height)

            if x < -0.01 or y < -0.01 or x + w > SW + 0.01 or y + h > SH + 0.01:
                hits.append(f"off-slide: {sh.shape_type}, "
                            f"x{x:.2f} y{y:.2f} w{w:.2f} h{h:.2f}")

            if not sh.has_text_frame:
                continue
            text = sh.text_frame.text
            if not text.strip():
                continue

            sizes = [r.font.size.pt for p in sh.text_frame.paragraphs
                     for r in p.runs if r.font.size]
            if not sizes:
                continue
            small = [s for s in sizes if s < MIN_PT]
            if small:
                hits.append(f"tiny text {min(small):.1f}pt: {text[:44]!r}")

            size = max(sizes)
            spacing = max((p.line_spacing or 1.2) for p in sh.text_frame.paragraphs)
            lines = est_lines(text, w, size)
            need = lines * size * spacing / 72.0
            if need > h + 0.14:
                hits.append(f"may overflow: needs ~{need:.2f}in, box {h:.2f}in "
                            f"({lines} lines @ {size:.0f}pt): {text[:40]!r}")

            boxes.append((x, y, w, h, text[:30]))

        # Overlapping text boxes, ignoring slivers.
        for i in range(len(boxes)):
            for j in range(i + 1, len(boxes)):
                ax, ay, aw, ah, at = boxes[i]
                bx, by, bw, bh, bt = boxes[j]
                ox = min(ax + aw, bx + bw) - max(ax, bx)
                oy = min(ay + ah, by + bh) - max(ay, by)
                if ox > 0.12 and oy > 0.12:
                    hits.append(f"text overlap {ox:.2f}x{oy:.2f}in: "
                                f"{at!r} / {bt!r}")

        if hits:
            problems += len(hits)
            print(f"slide {n}")
            for hit in hits:
                print(f"   {hit}")
            print()

    print("clean" if not problems else f"{problems} thing(s) to look at")


if __name__ == "__main__":
    main()
